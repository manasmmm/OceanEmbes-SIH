"""Extract the DrossVault MRV workbook into the dashboard's embedded data.

Usage:  python3 dashboard/build_data.py
Reads   data/DrossVault_Digital_MRV_Executive_Dashboard.xlsx
Writes  dashboard/index.html (from dashboard/src.html, replacing __DATA__)
"""
import json
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "DrossVault_Digital_MRV_Executive_Dashboard.xlsx"
SRC = ROOT / "dashboard" / "src.html"
OUT = ROOT / "dashboard" / "index.html"


def rows(ws, header_row=1):
    header = [c.value for c in ws[header_row]]
    out = []
    for r in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if all(v is None for v in r):
            continue
        out.append(dict(zip(header, r)))
    return out


def rnd(v, n=3):
    return None if v is None else round(float(v), n)


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=False)

    assumptions = [
        {"param": r["Parameter"], "value": r["Value"], "note": r["Notes"]}
        for r in rows(wb["Assumptions"])
    ]

    equipment = [
        {"name": r["Equipment"], "share": rnd(r["Share_%"], 1), "sec": rnd(r["Baseline_kWh_per_t"], 2)}
        for r in rows(wb["Equipment_Baseline"])
    ]

    daily = []
    for r in rows(wb["Daily_Summary"]):
        daily.append({
            "date": r["Date"].strftime("%Y-%m-%d"),
            "t": rnd(r["Throughput_t"]),
            "kwh": rnd(r["Electricity_kWh"], 1),
            "rec_t": rnd(r["Recovered_Material_t"]),
            "rej_t": rnd(r["Rejects_t"]),
            "down": rnd(r["Downtime_min"], 1),
            "run": rnd(r["Runtime_min"], 1),
            "co2": rnd(r["CO2e_kg"], 1),
            "sec": rnd(r["SEC_kWh_per_t"], 2),
            "rec": rnd(r["Recovery_Rate_pct"], 2),
            "co2t": rnd(r["CO2e_kg_per_t"], 2),
            "avail": rnd(r["Availability_pct"], 2),
        })

    opt = rows(wb["Optimization_Scenario"])
    for d, r in zip(daily, opt):
        d["sec_opt"] = rnd(r["Optimized_SEC_kWh_per_t"], 2)
        d["kwh_opt"] = rnd(r["Optimized_Electricity_kWh"], 1)
        d["saved"] = rnd(r["Energy_Saved_kWh"], 1)
        d["co2_avoided"] = rnd(r["CO2e_Avoided_kg"], 1)
        d["rec_opt"] = rnd(r["Optimized_Recovery_Rate_pct"], 2)

    kpi_ws = wb["Key_KPIs"]
    kpis = {}
    for r in kpi_ws.iter_rows(min_row=2, values_only=True):
        if r[0]:
            kpis[r[0]] = {"value": r[1], "unit": r[2]}

    # Hourly: compact column arrays
    state_codes = {"IDLE": 0, "STARTUP": 1, "RUNNING": 2, "SHUTDOWN": 3}
    hourly = {k: [] for k in ["t", "kwh", "sec", "rec", "moist", "temp", "amb", "curr", "vib",
                               "state", "eq", "en", "eqs", "ens", "diag"]}
    diag_list = []
    status_list = []
    for r in rows(wb["Hourly_MRV_Data"]):
        hourly["t"].append(rnd(r["Throughput_t"]))
        hourly["kwh"].append(rnd(r["Electricity_kWh"], 2))
        hourly["sec"].append(rnd(r["SEC_kWh_per_t"], 2))
        hourly["rec"].append(rnd(r["Recovery_Rate_pct"], 2))
        hourly["moist"].append(rnd(r["Moisture_pct"], 1))
        hourly["temp"].append(rnd(r["Temperature_C"], 1))
        hourly["amb"].append(rnd(r["Ambient_Temp_C"], 1))
        hourly["curr"].append(rnd(r["Motor_Current_A"], 1))
        hourly["vib"].append(rnd(r["Vibration_mm_s"], 2))
        hourly["state"].append(state_codes[r["Operating_State"]])
        hourly["eq"].append(r["Equipment_Anomaly_Score"])
        hourly["en"].append(r["Energy_Anomaly_Score"])
        for key, col in (("eqs", "Equipment_Anomaly_Status"), ("ens", "Energy_Anomaly_Status")):
            s = r[col]
            if s not in status_list:
                status_list.append(s)
            hourly[key].append(status_list.index(s))
        dg = r["Diagnosis_v2"]
        if dg not in diag_list:
            diag_list.append(dg)
        hourly["diag"].append(diag_list.index(dg))
    hourly["start"] = "2026-09-01T00:00"
    hourly["status_labels"] = status_list
    hourly["diag_labels"] = diag_list

    ac = wb["Anomaly_Config"]
    anomaly_cfg = []
    for r in ac.iter_rows(min_row=4, max_row=10, values_only=True):
        anomaly_cfg.append({"metric": r[0], "nl": r[1], "nh": r[2], "cl": r[3], "ch": r[4], "w": r[5]})

    alerts = []
    for r in rows(wb["Alert_Log"], header_row=3):
        alerts.append({
            "ts": r["Timestamp"].strftime("%Y-%m-%d %H:%M"),
            "state": r["Operating_State"],
            "sev": r["Severity"],
            "diag": r["Alert / Diagnosis"],
            "eq": r["Equipment_Score"],
            "en": r["Energy_Score"],
            "temp": r["Temperature_C"],
            "curr": r["Motor_Current_A"],
            "vib": r["Vibration_mm_s"],
            "tph": r["Throughput_tph"],
            "action": r["Recommended_Action"],
        })

    dq_rows = rows(wb["MRV_Data_Quality"], header_row=5)
    dq_rows = [r for r in dq_rows if r.get("Sensor_ID")]
    by_sensor = {}
    for r in dq_rows:
        s = by_sensor.setdefault(r["Sensor_ID"], {
            "sensor": r["Sensor_ID"], "equipment": r["Equipment_ID"], "measure": r["Measurement"],
            "unit": r["Unit"], "source": r["Source"],
            "last_cal": r["Last_Calibration"].strftime("%Y-%m-%d") if r["Last_Calibration"] else None,
            "valid": 0, "missing": 0, "verified": 0, "review": 0, "outliers": 0, "cal_valid": 0,
        })
        s["valid"] += r["Data_Status"] == "Valid"
        s["missing"] += r["Data_Status"] == "Missing"
        s["verified"] += r["Verification_Status"] == "Verified"
        s["review"] += r["Verification_Status"] == "Review"
        s["outliers"] += r["Outlier_Flag"] == "Yes"
        s["cal_valid"] += r["Calibration_Status"] == "Valid"
    dq = {
        "records": len(dq_rows),
        "hours": 720,
        "status": dict(Counter(r["Data_Status"] for r in dq_rows)),
        "verification": dict(Counter(r["Verification_Status"] for r in dq_rows)),
        "source": dict(Counter(r["Source"] for r in dq_rows)),
        "sensors": list(by_sensor.values()),
    }

    interventions = []
    for r in rows(wb["Optimization_Interventions"], header_row=5):
        if not r.get("Intervention") or r["Intervention"] == "TOTAL POTENTIAL":
            if r.get("Intervention") == "TOTAL POTENTIAL":
                break
            continue
        interventions.append({
            "name": r["Intervention"], "saving": r["SEC Saving (kWh/t)"], "evidence": r["Evidence / linkage"],
            "logic": r["Implementation Logic"], "priority": r["Priority"], "owner": r["Owner"],
            "kpi": r["Measurement KPI"], "status": r["Status"],
        })

    bc = wb["Business_Case"]
    business = {
        "tariff": bc["B5"].value, "capex": bc["B6"].value, "software": bc["B7"].value,
        "maintenance": bc["B8"].value, "days": bc["B9"].value, "sec_base": bc["B10"].value,
        "sec_opt": bc["B11"].value, "tpd": bc["B12"].value, "grid_ef": bc["B13"].value,
    }

    data = {
        "assumptions": assumptions, "equipment": equipment, "daily": daily, "kpis": kpis,
        "hourly": hourly, "anomaly_cfg": anomaly_cfg, "alerts": alerts, "dq": dq,
        "interventions": interventions, "business": business,
    }
    blob = json.dumps(data, separators=(",", ":"), ensure_ascii=False, default=str)
    html = SRC.read_text(encoding="utf-8").replace("__DATA__", blob)
    OUT.write_text(html, encoding="utf-8")
    print(f"wrote {OUT.relative_to(ROOT)} ({len(html)/1024:.0f} KB)")


if __name__ == "__main__":
    main()
